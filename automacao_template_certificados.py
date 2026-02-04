#pip install pandas python-pptx openpyxl
import pandas as pd
from datetime import datetime
df_topo = pd.read_excel("certificados_minicurso1.xlsx",header=None, nrows=1, sheet_name='Planilha2')

minicurso = str(df_topo.iloc[0,0])
Data_bruta = str(df_topo.iloc[0,1])
if isinstance(Data_bruta,(datetime, pd.Timestamp)):
    Data = Data_bruta.strftime('%d/%m/%Y')
else:
    Data = str(Data_bruta.split(' ')[0])
    if '-' in Data:
        partes = Data.split('-')
        if len(partes[0]) == 4:
            Data = f"{partes[2]}/{partes[1]}/{partes[0]}"


carga_horaria = str(df_topo.iloc[0,2])


df = pd.read_excel("certificados_minicurso1.xlsx", usecols="D,E", sheet_name='Planilha2')
df.columns = ["Nome", "Cpf"]
df['Cpf'] = df['Cpf'].astype(str).str.replace(r'\D','',regex=True)

def limpar_cpf(cpf):
    cpf = str(cpf).zfill(11)
    return f"{cpf[:3]}.{cpf[3:6]}.{cpf[6:9]}-{cpf[9:]}"

print(df.columns)
print()
print("-------------------")
print()
teste = df.head()
print(teste)
print()
print("-------------------")
print()

from pptx import Presentation

temmplate_padrao = Presentation("TEMPLATE_CERTIFICADO_QuIIN.pptx")
slide_modelo = temmplate_padrao.slides[0]

import copy

def copiar_slide_e_adicionar(temmplate_padrao, slide_modelo):
    slide_novo = temmplate_padrao.slides.add_slide(slide_modelo.slide_layout)

    for shape in list(slide_novo.shapes):
        if shape.is_placeholder:
            sp = shape._sp
            sp.getparent().remove(sp)

    for shape in slide_modelo.shapes:
        el = shape.element
        new_el = copy.deepcopy(el)
        
        slide_novo.shapes._spTree.insert_element_before(new_el, 'p:extLst')
    
    return slide_novo

def mudar_placeholder(paragraph, placeholders):
    texto = "".join(run.text for run in paragraph.runs)
    if not any(p in texto for p in placeholders.keys()):
        return
    
    paragraph.clear()
    
    i = 0
    while i< len(texto):
        achou = False
        for placeholder, valor in placeholders.items():
            if texto.startswith(placeholder,i):
                run = paragraph.add_run()
                run.text = valor["valor"]
                run.font.bold = valor["negrito"]
                i += len(placeholder)
                achou = True
                break
        if not achou:
            run = paragraph.add_run()
            run.text = texto[i]
            run.font.bold = False
            i +=1




for index, linha in df.iterrows():

    slide_atual = copiar_slide_e_adicionar(temmplate_padrao,slide_modelo)

    nome_atual = str(linha['Nome'])
    cpf_atual = limpar_cpf(linha['Cpf'])


    for shape in slide_atual.shapes:
        if shape.has_text_frame:
            for paragraph in shape.text_frame.paragraphs:
                placeholders = {
                    "[nome]":{"valor":nome_atual,"negrito":True},
                    "[cpf]":{"valor":cpf_atual, "negrito":True},
                    "[minicurso]":{"valor":minicurso,"negrito":True},
                    "[CH]":{"valor":carga_horaria,"negrito":True},
                    "[Data]":{"valor":Data,"negrito":False}
                }
                mudar_placeholder(paragraph,placeholders)

idSlide = temmplate_padrao.slides._sldIdLst[0].rId
temmplate_padrao.part.drop_rel(idSlide)
del temmplate_padrao.slides._sldIdLst[0]

    
temmplate_padrao.save("Certificados_unificados_1.pptx")
print("Todos os certificados foram gerados com sucesso!")
                
